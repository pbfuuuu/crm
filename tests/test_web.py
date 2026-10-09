"""Сквозные тесты веб-приложения (через тестовый клиент Flask)."""
import io
import os

from openpyxl import load_workbook

from tests.helpers import INN_A, INN_B, INN_C, TODAY, WebTestCase


class Auth(WebTestCase):
    def test_pages_require_login(self):
        c = self.app.test_client()
        for url in ("/", "/tenants/", "/premises/", "/contracts/", "/payments/", "/reports/debtors",
                    "/users/", "/audit/", "/payments/export", "/reports/debtors"):
            r = c.get(url)
            self.assertEqual(r.status_code, 302, url)
            self.assertIn("/login", r.headers["Location"], url)

    def test_login_success_and_failure_messages_are_generic(self):
        r = self.post(self.app.test_client(), "/login", {"username": "admin", "password": "AdminPass1"})
        self.assertEqual(r.status_code, 302)
        bad_pass = self.post(self.app.test_client(), "/login", {"username": "admin", "password": "nope"})
        bad_user = self.post(self.app.test_client(), "/login", {"username": "ghost", "password": "nope"})
        self.assertEqual(self.text(bad_pass), self.text(bad_user))   # не выдаём, есть ли такой логин
        self.assertIn("Неверный логин или пароль", self.text(bad_pass))

    def test_lockout_after_failed_attempts(self):
        c = self.app.test_client()
        for _ in range(5):
            self.post(c, "/login", {"username": "manager", "password": "wrong"})
        r = self.post(c, "/login", {"username": "manager", "password": "ManagerPass1"})
        self.assertIn("заблокирована", self.text(r))
        self.assertEqual(r.status_code, 200)               # даже верный пароль не пускает
        self.execute("UPDATE users SET locked_until = '2000-01-01T00:00:00' WHERE username='manager'")
        r = self.post(c, "/login", {"username": "manager", "password": "ManagerPass1"})
        self.assertEqual(r.status_code, 302)               # срок блокировки истёк

    def test_csrf_required(self):
        c = self.app.test_client()
        c.get("/login")
        r = c.post("/login", data={"username": "admin", "password": "AdminPass1"})
        self.assertEqual(r.status_code, 400)
        r = self.admin.post("/tenants/new", data={"name": "X", "inn": INN_A,
                                                  "csrf_token": "forged"})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self.one("SELECT COUNT(*) n FROM tenants")["n"], 0)

    def test_logout_is_post_only(self):
        self.assertEqual(self.admin.get("/logout").status_code, 405)
        r = self.post(self.admin, "/logout")
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.admin.get("/tenants/").status_code, 302)

    def test_open_redirect_blocked(self):
        for target in ("//evil.example", "https://evil.example", "javascript:alert(1)"):
            c = self.app.test_client()
            c.get("/login")
            with c.session_transaction() as s:
                s["_csrf"] = "t"
            r = c.post("/login?next=" + target, data={"username": "admin", "password": "AdminPass1",
                                                      "csrf_token": "t"})
            self.assertEqual(r.headers["Location"], "/", target)

    def test_local_next_is_honoured(self):
        c = self.app.test_client()
        c.get("/login")
        with c.session_transaction() as s:
            s["_csrf"] = "t"
        r = c.post("/login?next=/tenants/", data={"username": "admin", "password": "AdminPass1",
                                                  "csrf_token": "t"})
        self.assertEqual(r.headers["Location"], "/tenants/")

    def test_inactive_user_cannot_login_and_loses_session(self):
        self.execute("UPDATE users SET is_active = 0 WHERE username = 'manager'")
        self.assertEqual(self.manager.get("/tenants/").status_code, 302)  # сессия сброшена
        c = self.login("manager", "ManagerPass1")
        self.assertEqual(c.get("/tenants/").status_code, 302)

    def test_passwords_stored_as_hashes(self):
        for u in self.q("SELECT password_hash FROM users"):
            self.assertNotIn("AdminPass1", u["password_hash"])
            self.assertNotIn("ManagerPass1", u["password_hash"])
            self.assertTrue(u["password_hash"].startswith(("scrypt:", "pbkdf2:")))

    def test_security_headers(self):
        r = self.admin.get("/")
        self.assertIn("default-src 'self'", r.headers["Content-Security-Policy"])
        self.assertEqual(r.headers["X-Frame-Options"], "DENY")
        self.assertEqual(r.headers["X-Content-Type-Options"], "nosniff")
        self.assertEqual(r.headers["Cache-Control"], "no-store")
        login = self.post(self.app.test_client(), "/login", {"username": "admin", "password": "AdminPass1"})
        cookie = " ".join(login.headers.getlist("Set-Cookie"))
        self.assertIn("HttpOnly", cookie)
        self.assertIn("SameSite=Lax", cookie)

    def test_templates_have_no_inline_styles_or_scripts(self):
        """CSP запрещает inline — проверяем, что страницы не содержат style= и <script> без src."""
        t, p, cid = self.setup_contract()
        self.pay(cid, "2026-01", "50000")
        for url in ("/", "/tenants/", f"/tenants/{t}", "/premises/", "/contracts/",
                    f"/contracts/{cid}", "/payments/", "/payments/new", "/reports/debtors",
                    "/users/", "/audit/"):
            html = self.admin.get(url).get_data(as_text=True)
            self.assertNotIn(' style="', html, url)
            self.assertNotIn("onclick=", html, url)
            for tag in html.split("<script")[1:]:
                self.assertIn("src=", tag.split(">")[0], url)


class ForcedPasswordChange(WebTestCase):
    def setUp(self):
        super().setUp()
        self.execute("INSERT INTO users (username, full_name, password_hash, role, "
                     "must_change_password, created_at) VALUES ('temp','Temp',?, 'manager',1,'x')",
                     (self.hash("Temp0rary1"),))

    @staticmethod
    def hash(p):
        from werkzeug.security import generate_password_hash
        return generate_password_hash(p, method="pbkdf2:sha256:1000")

    def test_redirected_until_changed(self):
        c = self.login("temp", "Temp0rary1")
        for url in ("/", "/tenants/", "/payments/"):
            r = c.get(url)
            self.assertEqual(r.status_code, 302, url)
            self.assertIn("/password", r.headers["Location"], url)
        self.assertEqual(c.get("/password").status_code, 200)

    def test_weak_and_reused_passwords_rejected(self):
        c = self.login("temp", "Temp0rary1")
        for new, msg in (("short1", "не менее 8"), ("onlyletters", "буквы, и цифры"),
                         ("Temp0rary1", "отличаться")):
            r = self.post(c, "/password", {"old_password": "Temp0rary1", "new_password": new,
                                           "repeat_password": new})
            self.assertEqual(r.status_code, 200)
            self.assertIn(msg, self.text(r), new)
        r = self.post(c, "/password", {"old_password": "wrong", "new_password": "Brand1new9",
                                       "repeat_password": "Brand1new9"})
        self.assertIn("указан неверно", self.text(r))
        r = self.post(c, "/password", {"old_password": "Temp0rary1", "new_password": "Brand1new9",
                                       "repeat_password": "Different9"})
        self.assertIn("не совпадают", self.text(r))

    def test_change_unlocks_access(self):
        c = self.login("temp", "Temp0rary1")
        r = self.post(c, "/password", {"old_password": "Temp0rary1", "new_password": "Brand1new9",
                                       "repeat_password": "Brand1new9"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(c.get("/tenants/").status_code, 200)
        self.assertEqual(self.login("temp", "Temp0rary1").get("/tenants/").status_code, 302)
        self.assertEqual(self.login("temp", "Brand1new9").get("/tenants/").status_code, 200)
        self.assertEqual(self.one("SELECT must_change_password m FROM users WHERE username='temp'")["m"], 0)


class Roles(WebTestCase):
    def test_manager_forbidden_on_admin_pages(self):
        for url in ("/users/", "/users/new", "/audit/"):
            self.assertEqual(self.manager.get(url).status_code, 403, url)
            self.assertEqual(self.admin.get(url).status_code, 200, url)

    def test_manager_cannot_delete(self):
        t, p, cid = self.setup_contract()
        self.pay(cid, "2026-01", "50000")
        pid = self.one("SELECT id FROM payments")["id"]
        for url in (f"/tenants/{t}/delete", f"/premises/{p}/delete", f"/contracts/{cid}/delete",
                    f"/payments/{pid}/delete", "/users/1/reset", "/users/1/toggle"):
            self.assertEqual(self.post(self.manager, url).status_code, 403, url)
        self.assertEqual(self.one("SELECT COUNT(*) n FROM payments")["n"], 1)

    def test_manager_can_do_daily_work(self):
        t, p, cid = self.setup_contract()
        self.assertEqual(self.pay(cid, "2026-01", "50000").status_code, 302)

    def test_delete_buttons_hidden_for_manager(self):
        t, p, cid = self.setup_contract()
        self.assertNotIn("Удалить", self.text(self.manager.get(f"/tenants/{t}")))
        self.assertIn("Удалить", self.text(self.admin.get(f"/tenants/{t}")))
        self.assertNotIn("Пользователи", self.text(self.manager.get("/")))


class Tenants(WebTestCase):
    def test_create_encrypts_inn(self):
        tid = self.add_tenant(inn=INN_A)
        row = self.one("SELECT * FROM tenants WHERE id = ?", (tid,))
        self.assertNotIn(INN_A, str(row))
        self.assertNotEqual(row["inn_encrypted"], INN_A)
        self.assertEqual(len(row["inn_hash"]), 64)
        self.assertIn(INN_A, self.text(self.manager.get(f"/tenants/{tid}")))   # в интерфейсе виден
        # в файлах базы ИНН в открытом виде отсутствует
        for suffix in ("", "-wal"):
            path = self.db_path + suffix
            if os.path.exists(path):
                with open(path, "rb") as f:
                    self.assertNotIn(INN_A.encode(), f.read(), path)

    def test_validation(self):
        c = self.manager
        r = self.post(c, "/tenants/new", {"name": "", "inn": "123"})
        t = self.text(r)
        self.assertIn("обязательно", t)
        self.assertIn("ИНН указан неверно", t)
        r = self.post(c, "/tenants/new", {"name": "Тест", "inn": INN_A, "email": "not-mail"})
        self.assertIn("корректный адрес", self.text(r))
        r = self.post(c, "/tenants/new", {"name": "Тест", "inn": INN_A, "phone": "12"})
        self.assertIn("от 10 до 15 цифр", self.text(r))
        self.assertEqual(self.one("SELECT COUNT(*) n FROM tenants")["n"], 0)

    def test_duplicate_inn_rejected(self):
        self.add_tenant("Первый", INN_A)
        r = self.post(self.manager, "/tenants/new", {"name": "Второй", "inn": INN_A})
        self.assertIn("уже существует", self.text(r))
        r = self.post(self.manager, "/tenants/new", {"name": "Второй", "inn": " 7707 083893 "})
        self.assertIn("уже существует", self.text(r))        # пробелы не помогают обойти

    def test_edit_changes_data_and_inn(self):
        tid = self.add_tenant("Старое", INN_A)
        self.add_tenant("Другой", INN_B)
        r = self.post(self.manager, f"/tenants/{tid}/edit", {"name": "Новое", "inn": INN_B})
        self.assertIn("уже существует", self.text(r))
        r = self.post(self.manager, f"/tenants/{tid}/edit", {"name": "Новое", "inn": INN_C})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.one("SELECT name FROM tenants WHERE id=?", (tid,))["name"], "Новое")
        self.assertIn(INN_C, self.text(self.manager.get(f"/tenants/{tid}")))
        self.assertIn("ИНН", self.one("SELECT details FROM audit_log WHERE action='update'")["details"])
        self.assertNotIn(INN_C, str(self.q("SELECT * FROM audit_log")))   # ИНН не попадает в журнал

    def test_search_is_case_insensitive_cyrillic(self):
        self.add_tenant("ООО «Ромашка»", INN_A)
        self.add_tenant("ИП Смирнов", INN_B)
        t = self.text(self.manager.get("/tenants/?q=РОМАШКА"))
        self.assertIn("Ромашка", t)
        self.assertNotIn("Смирнов", t)
        t = self.text(self.manager.get("/tenants/?q=смирнов"))
        self.assertIn("Смирнов", t)

    def test_delete_blocked_with_contract(self):
        t, p, cid = self.setup_contract()
        r = self.post(self.admin, f"/tenants/{t}/delete", follow=True)
        self.assertIn("есть договоры", self.text(r))
        self.assertEqual(self.one("SELECT COUNT(*) n FROM tenants")["n"], 1)
        t2 = self.add_tenant("Свободный", INN_B)
        self.post(self.admin, f"/tenants/{t2}/delete")
        self.assertIsNone(self.one("SELECT id FROM tenants WHERE id=?", (t2,)))

    def test_html_is_escaped(self):
        tid = self.add_tenant("<script>alert(1)</script>", INN_A)
        html = self.manager.get(f"/tenants/{tid}").get_data(as_text=True)
        self.assertNotIn("<script>alert(1)</script>", html)
        self.assertIn("&lt;script&gt;", html)

    def test_wrong_key_does_not_crash(self):
        tid = self.add_tenant()
        self.execute("UPDATE tenants SET inn_encrypted = 'garbage' WHERE id = ?", (tid,))
        r = self.manager.get(f"/tenants/{tid}")
        self.assertEqual(r.status_code, 200)
        self.assertIn("не удалось расшифровать", self.text(r))


class Premises(WebTestCase):
    def test_create_validation_and_duplicates(self):
        self.add_premises("101")
        r = self.post(self.manager, "/premises/new", {"number": "101", "floor": "1", "area": "10", "status": "free"})
        self.assertIn("уже существует", self.text(r))
        for data in ({"number": "X", "floor": "abc", "area": "10"},
                     {"number": "X", "floor": "1", "area": "-5"},
                     {"number": "X", "floor": "1", "area": "0"},
                     {"number": "X", "floor": "1", "area": "10", "status": "rented"}):
            data.setdefault("status", "free")
            r = self.post(self.manager, "/premises/new", data)
            self.assertEqual(r.status_code, 200, data)
        self.assertEqual(self.one("SELECT COUNT(*) n FROM premises")["n"], 1)

    def test_decimal_comma_area(self):
        self.add_premises("102", area="45,5")
        self.assertEqual(self.one("SELECT area FROM premises WHERE number='102'")["area"], 45.5)

    def test_status_locked_when_rented(self):
        t, p, cid = self.setup_contract()
        self.assertIn('value="Занято (есть действующий договор)"',
                      self.manager.get(f"/premises/{p}/edit").get_data(as_text=True))
        self.post(self.manager, f"/premises/{p}/edit", {"number": "101", "floor": "1", "area": "50",
                                                        "status": "repair"})
        self.assertEqual(self.one("SELECT status FROM premises WHERE id=?", (p,))["status"], "rented")

    def test_delete_blocked_with_contracts(self):
        t, p, cid = self.setup_contract()
        r = self.post(self.admin, f"/premises/{p}/delete", follow=True)
        self.assertIn("есть договоры", self.text(r))
        self.assertEqual(self.one("SELECT COUNT(*) n FROM premises")["n"], 1)

    def test_filters(self):
        self.add_premises("101", "1"); self.add_premises("201", "2")
        self.assertNotIn("201", self.text(self.manager.get("/premises/?floor=1")))
        self.assertIn("201", self.text(self.manager.get("/premises/?floor=2")))


class Contracts(WebTestCase):
    def test_create_marks_premises_rented(self):
        t, p, cid = self.setup_contract()
        self.assertEqual(self.one("SELECT status FROM premises WHERE id=?", (p,))["status"], "rented")
        self.assertEqual(self.one("SELECT monthly_rent r FROM contracts")["r"], 5000000)
        r = self.manager.get("/contracts/new")
        self.assertNotIn("101 · 1 эт.", self.text(r))        # занятое помещение не предлагается

    def test_validation(self):
        t = self.add_tenant(); p = self.add_premises()
        base = {"tenant_id": t, "premises_id": p, "number": "Д-1", "start_date": "2026-01-01",
                "end_date": "2026-12-31", "monthly_rent": "50000"}
        cases = [({"end_date": "2025-12-31"}, "раньше даты начала"),
                 ({"monthly_rent": "0"}, "больше нуля"), ({"monthly_rent": "abc"}, "больше нуля"),
                 ({"monthly_rent": "999999999999"}, "больше нуля"),
                 ({"start_date": ""}, "дату начала"), ({"tenant_id": ""}, "Выберите клиента"),
                 ({"premises_id": "999"}, "Выберите помещение"), ({"number": ""}, "обязательно")]
        for patch, msg in cases:
            r = self.post(self.manager, "/contracts/new", {**base, **patch})
            self.assertEqual(r.status_code, 200, patch)
            self.assertIn(msg, self.text(r), patch)
        self.assertEqual(self.one("SELECT COUNT(*) n FROM contracts")["n"], 0)

    def test_cannot_rent_occupied_or_repair_premises(self):
        t, p, cid = self.setup_contract()
        t2 = self.add_tenant("Другой", INN_B)
        r = self.post(self.manager, "/contracts/new", {
            "tenant_id": t2, "premises_id": p, "number": "Д-2", "start_date": "2026-01-01",
            "end_date": "2026-12-31", "monthly_rent": "1000"})
        self.assertIn("недоступно для аренды", self.text(r))
        p2 = self.add_premises("102")
        self.execute("UPDATE premises SET status='repair' WHERE id=?", (p2,))
        r = self.post(self.manager, "/contracts/new", {
            "tenant_id": t2, "premises_id": p2, "number": "Д-3", "start_date": "2026-01-01",
            "end_date": "2026-12-31", "monthly_rent": "1000"})
        self.assertIn("недоступно для аренды", self.text(r))

    def test_overlap_detected_even_if_status_wrong(self):
        t, p, cid = self.setup_contract(start="2026-01-01", end="2026-12-31")
        self.execute("UPDATE premises SET status='free' WHERE id=?", (p,))   # рассинхрон статуса
        t2 = self.add_tenant("Другой", INN_B)
        r = self.post(self.manager, "/contracts/new", {
            "tenant_id": t2, "premises_id": p, "number": "Д-2", "start_date": "2026-06-01",
            "end_date": "2027-06-01", "monthly_rent": "1000"})
        self.assertIn("уже есть действующий договор", self.text(r))

    def test_duplicate_number(self):
        t, p, cid = self.setup_contract(number="Д-1")
        t2 = self.add_tenant("Другой", INN_B); p2 = self.add_premises("102")
        r = self.post(self.manager, "/contracts/new", {
            "tenant_id": t2, "premises_id": p2, "number": "Д-1", "start_date": "2026-01-01",
            "end_date": "2026-12-31", "monthly_rent": "1000"})
        self.assertIn("уже существует", self.text(r))

    def test_debt_grows_with_calendar(self):
        """Договор с 01.01.2026, сегодня 06.10.2026: 10 месяцев × 50 000 = 500 000."""
        t, p, cid = self.setup_contract()
        t_ = self.text(self.manager.get(f"/contracts/{cid}"))
        self.assertIn("500 000,00 ₽", t_)

    def test_rent_change_only_affects_future_months(self):
        t, p, cid = self.setup_contract()
        r = self.post(self.manager, f"/contracts/{cid}/rent", {"new_rent": "55000", "valid_from": "2026-04"}, follow=True)
        self.assertIn("Ранее начисленные суммы не изменились", self.text(r))
        # янв–март по 50 000, апр–окт (7 мес.) по 55 000
        self.assertIn(f"{3 * 50000 + 7 * 55000:,}".replace(",", " ") + ",00 ₽", self.text(r))
        with self.app.app_context():
            from app import services
            from app.db import get_db
            data = services.load_calc_data(get_db())
            c = dict(get_db().execute("SELECT * FROM contracts").fetchone())
            self.assertEqual(services.rent_for_month(c, data, "2026-03"), 5000000)
            self.assertEqual(services.rent_for_month(c, data, "2026-04"), 5500000)

    def test_rent_change_validation(self):
        t, p, cid = self.setup_contract()
        for data, msg in (({"new_rent": "0", "valid_from": "2026-04"}, "больше нуля"),
                          ({"new_rent": "100", "valid_from": "2027-05"}, "должен попадать"),
                          ({"new_rent": "100", "valid_from": "2025-12"}, "должен попадать"),
                          ({"new_rent": "100", "valid_from": "bad"}, "должен попадать")):
            r = self.post(self.manager, f"/contracts/{cid}/rent", data, follow=True)
            self.assertIn(msg, self.text(r), data)
        self.post(self.manager, f"/contracts/{cid}/rent", {"new_rent": "60000", "valid_from": "2026-05"})
        r = self.post(self.manager, f"/contracts/{cid}/rent", {"new_rent": "70000", "valid_from": "2026-05"}, follow=True)
        self.assertIn("уже изменялась", self.text(r))
        self.assertEqual(self.one("SELECT COUNT(*) n FROM rent_changes")["n"], 1)

    def test_only_admin_deletes_rent_change(self):
        t, p, cid = self.setup_contract()
        self.post(self.manager, f"/contracts/{cid}/rent", {"new_rent": "60000", "valid_from": "2026-05"})
        rid = self.one("SELECT id FROM rent_changes")["id"]
        self.assertEqual(self.post(self.manager, f"/contracts/{cid}/rent/{rid}/delete").status_code, 403)
        self.post(self.admin, f"/contracts/{cid}/rent/{rid}/delete")
        self.assertEqual(self.one("SELECT COUNT(*) n FROM rent_changes")["n"], 0)

    def test_terminate_frees_premises_and_stops_accrual(self):
        t, p, cid = self.setup_contract()
        r = self.post(self.manager, f"/contracts/{cid}/terminate", {"end_date": "2026-06-30"}, follow=True)
        self.assertIn("Договор закрыт", self.text(r))
        row = self.one("SELECT * FROM contracts WHERE id=?", (cid,))
        self.assertEqual((row["is_active"], row["end_date"]), (0, "2026-06-30"))
        self.assertEqual(self.one("SELECT status FROM premises WHERE id=?", (p,))["status"], "free")
        self.assertIn("300 000,00 ₽", self.text(self.manager.get(f"/contracts/{cid}")))   # 6 мес.
        r = self.post(self.manager, f"/contracts/{cid}/terminate", {"end_date": "2026-06-30"}, follow=True)
        self.assertIn("уже закрыт", self.text(r))
        r = self.post(self.manager, f"/contracts/{cid}/rent", {"new_rent": "1", "valid_from": "2026-02"}, follow=True)
        self.assertIn("закрытого договора", self.text(r))
        self.assertIn("Сдать", self.text(self.manager.get("/premises/")))     # снова доступно к сдаче
        # помещение снова можно сдать
        t2 = self.add_tenant("Новый", INN_B)
        self.add_contract(t2, p, number="Д-2", start="2026-07-01")

    def test_terminate_date_must_be_in_term(self):
        t, p, cid = self.setup_contract()
        for bad in ("2025-12-31", "2027-01-01", "", "xx"):
            r = self.post(self.manager, f"/contracts/{cid}/terminate", {"end_date": bad}, follow=True)
            self.assertIn("в пределах срока", self.text(r), bad)
        self.assertEqual(self.one("SELECT is_active a FROM contracts")["a"], 1)

    def test_expired_contract_flagged(self):
        t, p, cid = self.setup_contract(start="2025-01-01", end="2026-09-30")
        t_ = self.text(self.admin.get("/"))
        self.assertIn("Срок истёк, договор не закрыт", t_)
        self.assertIn("Д-1", t_)
        self.assertIn("Д-1", self.text(self.manager.get("/contracts/?state=expired")))
        self.assertNotIn("Д-1", self.text(self.manager.get("/contracts/?state=active")))

    def test_ending_soon_filter(self):
        self.setup_contract(end="2026-11-15")      # через 40 дней — попадает в 60 дней
        t2 = self.add_tenant("Другой", INN_B); p2 = self.add_premises("102")
        self.add_contract(t2, p2, number="Д-2", end="2027-12-31")
        t = self.text(self.manager.get("/contracts/?state=ending"))
        self.assertIn("Д-1", t)
        self.assertNotIn("Д-2", t)

    def test_edit_end_date(self):
        t, p, cid = self.setup_contract()
        r = self.post(self.manager, f"/contracts/{cid}/edit", {"number": "Д-1", "end_date": "2027-06-30"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.one("SELECT end_date e FROM contracts")["e"], "2027-06-30")
        r = self.post(self.manager, f"/contracts/{cid}/edit", {"number": "Д-1", "end_date": "2025-01-01"})
        self.assertIn("раньше даты начала", self.text(r))

    def test_delete_blocked_with_payments(self):
        t, p, cid = self.setup_contract()
        self.pay(cid, "2026-01", "50000")
        r = self.post(self.admin, f"/contracts/{cid}/delete", follow=True)
        self.assertIn("есть оплаты", self.text(r))
        self.assertEqual(self.one("SELECT COUNT(*) n FROM contracts")["n"], 1)

    def test_delete_without_payments_frees_premises(self):
        t, p, cid = self.setup_contract()
        self.post(self.admin, f"/contracts/{cid}/delete")
        self.assertEqual(self.one("SELECT COUNT(*) n FROM contracts")["n"], 0)
        self.assertEqual(self.one("SELECT status FROM premises WHERE id=?", (p,))["status"], "free")


class Payments(WebTestCase):
    def test_payment_reduces_debt(self):
        t, p, cid = self.setup_contract()
        self.assertIn("500 000,00 ₽", self.text(self.manager.get(f"/contracts/{cid}")))
        r = self.pay(cid, "2026-01", "50000")
        self.assertEqual(r.status_code, 302)
        t_ = self.text(self.manager.get(f"/contracts/{cid}"))
        self.assertIn("450 000,00 ₽", t_)

    def test_partial_payment_status(self):
        t, p, cid = self.setup_contract()
        self.pay(cid, "2026-10", "20000")
        t_ = self.text(self.manager.get("/payments/?month=2026-10"))
        self.assertIn("Частично", t_)
        self.assertIn("30 000,00 ₽", t_)                       # остаток

    def test_month_view_and_unpaid_filter(self):
        t, p, cid = self.setup_contract()
        t2 = self.add_tenant("Плательщик", INN_B); p2 = self.add_premises("102")
        c2 = self.add_contract(t2, p2, number="Д-2", rent="10000")
        self.pay(c2, "2026-10", "10000")
        all_ = self.text(self.manager.get("/payments/?month=2026-10"))
        self.assertIn("Ромашка", all_); self.assertIn("Плательщик", all_)
        unpaid = self.text(self.manager.get("/payments/?month=2026-10&filter=unpaid"))
        self.assertIn("Ромашка", unpaid); self.assertNotIn("Плательщик", unpaid)

    def test_future_month_not_due(self):
        t, p, cid = self.setup_contract()
        t_ = self.text(self.manager.get("/payments/?month=2026-12"))
        self.assertIn("Не наступил", t_)
        self.assertNotIn("Не оплачено", t_)

    def test_validation(self):
        t, p, cid = self.setup_contract(start="2026-03-01", end="2026-12-31")
        cases = [(("2026-02", "100", "2026-10-05"), "не входит в срок"),
                 (("2026-11", "100", "2026-10-05"), "будущий месяц"),
                 (("2026-10", "0", "2026-10-05"), "больше нуля"),
                 (("2026-10", "abc", "2026-10-05"), "больше нуля"),
                 (("2026-10", "100", "2026-10-07"), "в будущем"),
                 (("2026-10", "100", ""), "дату оплаты"),
                 (("", "100", "2026-10-05"), "месяц")]
        for (period, amount, paid_at), msg in cases:
            r = self.pay(cid, period, amount, paid_at)
            self.assertEqual(r.status_code, 200, (period, amount, paid_at))
            self.assertIn(msg, self.text(r), (period, amount, paid_at))
        r = self.pay(9999, "2026-10", "100")
        self.assertIn("Выберите договор", self.text(r))
        self.assertEqual(self.one("SELECT COUNT(*) n FROM payments")["n"], 0)

    def test_overpayment_is_recorded_as_overpaid(self):
        t, p, cid = self.setup_contract(start="2026-10-01")
        self.pay(cid, "2026-10", "80000")
        t_ = self.text(self.manager.get(f"/contracts/{cid}"))
        self.assertIn("Переплата", t_)
        self.assertIn("30 000,00 ₽", t_)

    def test_prefilled_amount_is_remaining(self):
        t, p, cid = self.setup_contract(start="2026-10-01")
        self.pay(cid, "2026-10", "20000")
        html = self.manager.get(f"/payments/new?contract_id={cid}&period=2026-10").get_data(as_text=True)
        self.assertIn('value="30000.00"', html)

    def test_payment_on_closed_contract_allowed(self):
        t, p, cid = self.setup_contract()
        self.post(self.manager, f"/contracts/{cid}/terminate", {"end_date": "2026-06-30"})
        self.assertEqual(self.pay(cid, "2026-05", "50000").status_code, 302)
        self.assertEqual(self.pay(cid, "2026-07", "50000").status_code, 200)   # после расторжения - нет

    def test_next_redirect_is_local_only(self):
        t, p, cid = self.setup_contract()
        r = self.post(self.manager, "/payments/new", {
            "contract_id": cid, "period": "2026-01", "amount": "100", "paid_at": "2026-10-05",
            "next": "https://evil.example/"})
        self.assertTrue(r.headers["Location"].startswith("/payments/"))
        r = self.post(self.manager, "/payments/new", {
            "contract_id": cid, "period": "2026-02", "amount": "100", "paid_at": "2026-10-05",
            "next": f"/contracts/{cid}"})
        self.assertEqual(r.headers["Location"], f"/contracts/{cid}")

    def test_admin_deletes_payment(self):
        t, p, cid = self.setup_contract()
        self.pay(cid, "2026-01", "50000")
        pid = self.one("SELECT id FROM payments")["id"]
        self.post(self.admin, f"/payments/{pid}/delete")
        self.assertEqual(self.one("SELECT COUNT(*) n FROM payments")["n"], 0)
        self.assertEqual(self.one("SELECT action FROM audit_log WHERE action='payment_delete'")["action"], "payment_delete")


class Reports(WebTestCase):
    def setUp(self):
        super().setUp()
        self.t1, self.p1, self.c1 = self.setup_contract(rent="50000")           # долг с января
        self.t2 = self.add_tenant("Аккуратный", INN_B); self.p2 = self.add_premises("102", area="100")
        self.c2 = self.add_contract(self.t2, self.p2, number="Д-2", start="2026-09-01", rent="10000")
        self.pay(self.c2, "2026-09", "10000"); self.pay(self.c2, "2026-10", "10000")
        self.add_premises("103", area="50")          # свободное

    def test_debtors_report(self):
        t = self.text(self.manager.get("/reports/debtors"))
        self.assertIn("Ромашка", t)
        self.assertNotIn("Аккуратный", t)
        self.assertIn("500 000,00 ₽", t)

    def test_debtors_excel(self):
        r = self.manager.get("/reports/debtors?format=xlsx")
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.data)).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "Клиент")
        self.assertEqual(rows[1][0], "ООО «Ромашка»")
        self.assertEqual(rows[1][6], 500000)           # рубли, а не копейки
        self.assertEqual(rows[-1][0], "ИТОГО")
        self.assertEqual(rows[-1][6], 500000)

    def test_month_excel(self):
        r = self.manager.get("/payments/export?month=2026-10")
        ws = load_workbook(io.BytesIO(r.data)).active
        rows = list(ws.iter_rows(values_only=True))
        body = {row[0]: row for row in rows[1:-1]}
        self.assertEqual(body["Д-1"][3:6], (50000, 0, 50000))
        self.assertEqual(body["Д-2"][3:6], (10000, 10000, 0))
        self.assertEqual(rows[-1][3:6], (60000, 10000, 50000))

    def test_dashboard_numbers(self):
        t = self.text(self.admin.get("/"))
        self.assertIn("75%", t)
        self.assertIn("занято 2 из 3 помещений", t)
        self.assertIn("начислено 60 000,00 ₽", t)       # 50 000 + 10 000 за октябрь
        self.assertIn("10 000,00 ₽", t)                 # получено
        self.assertIn("500 000,00 ₽", t)                # общий долг

    def test_exports_are_logged(self):
        self.manager.get("/reports/debtors?format=xlsx")
        self.assertEqual(self.one("SELECT COUNT(*) n FROM audit_log WHERE action='export'")["n"], 1)


class Users(WebTestCase):
    def test_create_user_shows_password_once_and_forces_change(self):
        r = self.post(self.admin, "/users/new", {"username": "Ivan", "full_name": "Иван", "role": "manager"})
        html = r.get_data(as_text=True)
        self.assertEqual(r.status_code, 200)
        import html as htmllib, re
        pwd = htmllib.unescape(re.search(r'class="password-box">([^<]+)<', html).group(1))
        self.assertEqual(len(pwd), 12)
        row = self.one("SELECT * FROM users WHERE username='ivan'")       # логин приведён к нижнему регистру
        self.assertEqual(row["must_change_password"], 1)
        self.assertNotIn(pwd, str(row))
        self.assertNotIn(pwd, str(self.q("SELECT * FROM audit_log")))
        c = self.login("ivan", pwd)
        self.assertIn("/password", c.get("/tenants/").headers["Location"])

    def test_username_validation(self):
        for bad in ("ab", "with space", "Русский", "a" * 31, "bad/char"):
            r = self.post(self.admin, "/users/new", {"username": bad, "full_name": "X", "role": "manager"})
            self.assertIn("Логин: 3–30", self.text(r), bad)
        r = self.post(self.admin, "/users/new", {"username": "ADMIN", "full_name": "X", "role": "manager"})
        self.assertIn("уже существует", self.text(r))
        r = self.post(self.admin, "/users/new", {"username": "okname", "full_name": "X", "role": "root"})
        self.assertIn("Выберите роль", self.text(r))

    def test_reset_password(self):
        uid = self.one("SELECT id FROM users WHERE username='manager'")["id"]
        r = self.post(self.admin, f"/users/{uid}/reset")
        import html as htmllib, re
        pwd = htmllib.unescape(re.search(r'class="password-box">([^<]+)<', r.get_data(as_text=True)).group(1))
        self.assertEqual(self.login("manager", "ManagerPass1").get("/tenants/").status_code, 302)
        c = self.login("manager", pwd)
        self.assertIn("/password", c.get("/tenants/").headers["Location"])

    def test_cannot_lock_self(self):
        uid = self.one("SELECT id FROM users WHERE username='admin'")["id"]
        r = self.post(self.admin, f"/users/{uid}/toggle", follow=True)
        self.assertIn("собственную", self.text(r))
        self.assertEqual(self.one("SELECT is_active a FROM users WHERE id=?", (uid,))["a"], 1)

    def test_toggle_blocks_login(self):
        uid = self.one("SELECT id FROM users WHERE username='manager'")["id"]
        self.post(self.admin, f"/users/{uid}/toggle")
        self.assertEqual(self.login("manager", "ManagerPass1").get("/tenants/").status_code, 302)
        self.post(self.admin, f"/users/{uid}/toggle")
        self.assertEqual(self.login("manager", "ManagerPass1").get("/tenants/").status_code, 200)


class AuditLog(WebTestCase):
    def test_actions_are_logged(self):
        t, p, cid = self.setup_contract()
        self.pay(cid, "2026-01", "50000")
        self.post(self.manager, f"/contracts/{cid}/rent", {"new_rent": "60000", "valid_from": "2026-05"})
        actions = {r["action"] for r in self.q("SELECT action FROM audit_log")}
        self.assertTrue({"login", "create", "payment_add", "rent_change"} <= actions)
        row = self.one("SELECT * FROM audit_log WHERE action='payment_add'")
        self.assertEqual(row["username"], "manager")
        self.assertIn("50 000,00 ₽", row["details"])

    def test_failed_login_logged_without_password(self):
        self.post(self.app.test_client(), "/login", {"username": "admin", "password": "SecretGuess1"})
        rows = self.q("SELECT * FROM audit_log WHERE action='login_failed'")
        self.assertEqual(len(rows), 1)
        self.assertNotIn("SecretGuess1", str(rows))

    def test_view_and_filters(self):
        self.setup_contract()
        html = self.admin.get("/audit/?action=create").get_data(as_text=True)
        n_create = self.one("SELECT COUNT(*) n FROM audit_log WHERE action='create'")["n"]
        self.assertGreater(n_create, 0)
        self.assertEqual(html.count("<tr>") - 1, n_create)      # только события «Создание»
        t = self.text(self.admin.get("/audit/?user=manager"))
        self.assertIn("manager", t)


class Misc(WebTestCase):
    def test_404_page(self):
        r = self.admin.get("/tenants/9999")
        self.assertEqual(r.status_code, 404)
        self.assertIn("Страница не найдена", self.text(r))

    def test_sql_injection_attempts_are_harmless(self):
        self.add_tenant()
        for payload in ("' OR '1'='1", "'; DROP TABLE tenants; --", "%", "_"):
            self.assertEqual(self.manager.get("/tenants/", query_string={"q": payload}).status_code, 200)
            self.assertEqual(self.admin.get("/audit/", query_string={"user": payload}).status_code, 200)
        self.assertEqual(self.one("SELECT COUNT(*) n FROM tenants")["n"], 1)

    def test_backup(self):
        import sqlite3
        self.add_tenant()
        dst = os.path.join(self.tmp.name, "copy.sqlite3")
        s, d = sqlite3.connect(self.db_path), sqlite3.connect(dst)
        s.backup(d); s.close(); d.close()
        self.assertEqual(sqlite3.connect(dst).execute("SELECT COUNT(*) FROM tenants").fetchone()[0], 1)


class Seed(WebTestCase):
    def test_seed_produces_expected_demo_numbers(self):
        from seed import seed
        # чистая вторая БД, чтобы не смешивать с пользователями из setUp
        import tempfile
        from cryptography.fernet import Fernet
        from app import create_app
        with tempfile.TemporaryDirectory() as d:
            app2 = create_app({"DATABASE": os.path.join(d, "s.sqlite3"), "SECRET_KEY": "k",
                               "FERNET_KEY": Fernet.generate_key().decode(), "TESTING": True})
            seed(app2)
            with app2.app_context():
                from app import services
                from app.db import get_db
                contracts, _ = services.contracts_with_summaries(get_db())
                debts = {c["number"]: c["summary"]["debt"] for c in contracts}
                self.assertEqual(debts["А-021/25"], 12000000)
                self.assertEqual(debts["А-018/25"], 7500000)
                self.assertEqual(debts["А-030/26"], 2400000)
                self.assertEqual(sum(debts.values()), 21900000)
                states = {c["number"]: c["state"] for c in contracts}
                self.assertEqual(states["А-018/25"], "expired")
                self.assertEqual(states["А-005/25"], "closed")


if __name__ == "__main__":
    import unittest
    unittest.main()
