"""Экспорт отчётов в Excel."""
import io
from datetime import datetime

from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def xlsx_response(filename, sheet_title, headers, rows, money_cols=(), total_row=None):
    """money_cols - индексы колонок, где значения в копейках (будут переведены в рубли)."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E2F3")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows + ([total_row] if total_row else []):
        ws.append([(v / 100 if i in money_cols and isinstance(v, (int, float)) else v)
                   for i, v in enumerate(row)])
    for i in money_cols:
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=i + 1).number_format = '#,##0.00'
    if total_row:
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(ws.cell(row=r, column=i).value or ""))
                                     for r in range(2, ws.max_row + 1)])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 3, 50)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = datetime.now().strftime("%Y%m%d")
    return send_file(buf, as_attachment=True, download_name=f"{filename}_{stamp}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
